# -*- coding: utf-8 -*-
"""油补里程 × 任务完成度 / 住宿费扣减分析（v2）

确认口径：
1A 无对应审核通过任务 → 扣当天有效里程
2A 完成度不达标 → 扣当天有效里程（同人同日取最高完成度）
3 任务有效：创建/修改审核通过
有申诉即免扣：覆盖该天任一条「已申诉」任务 → 里程不扣
鸡窝：仅 >=2026-08-24；大宗：2026-08 ~ 2026-09
住宿费：实际>0用实际，否则用申请；未达标/无有效任务且未申诉则扣；已申诉不扣
统计：分母=期内有有效里程的人；没创建任务=至少1天无审核通过任务
"""
from __future__ import annotations

from collections import defaultdict
from datetime import datetime, date
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

SRC = Path(r"c:\Users\HY\Desktop\cursor分析数据原始表.xlsx")
OUT = Path(r"c:\Users\HY\Desktop\油补里程任务扣减分析_v2.xlsx")
OUT2 = Path(r"d:\trae\crm-task\油补里程任务扣减分析.xlsx")
OUT3 = Path(r"c:\Users\HY\Desktop\油补里程任务扣减分析.xlsx")

PASS_STATUS = {"创建审核通过", "修改审核通过"}
THRESH = {"大宗": 1.0, "鸡窝": 0.75}
JIWO_START = date(2026, 8, 24)
DAZONG_START = date(2026, 8, 1)
DAZONG_END = date(2026, 9, 30)

# 里程表 / 任务表姓名不一致
NAME_ALIASES = {
    "张桐?": "张桐菓",
    "王志刚1": "王志刚",
}


def parse_pct(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        x = float(v)
        return x / 100.0 if x > 1 else x
    s = str(v).strip().replace("%", "")
    try:
        return float(s) / 100.0
    except Exception:
        return None


def parse_day(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()[:10]
    for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except Exception:
            continue
    return None


def parse_range(s):
    if s is None:
        return None, None
    text = str(s).strip()
    if " ~ " in text:
        a, b = text.split(" ~ ", 1)
    elif "~" in text:
        a, b = text.split("~", 1)
    else:
        d = parse_day(text)
        return d, d
    return parse_day(a.strip()), parse_day(b.strip())


def in_window(biz, day):
    if biz == "鸡窝":
        return day >= JIWO_START
    if biz == "大宗":
        return DAZONG_START <= day <= DAZONG_END
    # 未知：按大宗窗口（8-9月），且若 <8.24 仍计入（大宗可能）
    return DAZONG_START <= day <= DAZONG_END


def norm_name(name: str) -> str:
    name = (name or "").strip()
    return NAME_ALIASES.get(name, name)


def lodging_fee(apply_fee, actual_fee):
    try:
        actual = float(actual_fee or 0)
    except Exception:
        actual = 0.0
    try:
        apply = float(apply_fee or 0)
    except Exception:
        apply = 0.0
    return actual if actual > 0 else apply


def load_tasks(ws):
    tasks = []
    for r in range(2, ws.max_row + 1):
        name = norm_name(str(ws.cell(r, 4).value or "").strip())
        biz = str(ws.cell(r, 5).value or "").strip()
        if not name:
            continue
        start, end = parse_range(ws.cell(r, 9).value)
        if not start or not end:
            continue
        if end < start:
            start, end = end, start
        status = str(ws.cell(r, 3).value or "").strip()
        pct = parse_pct(ws.cell(r, 8).value)
        appeal = str(ws.cell(r, 10).value or "").strip() == "已申诉"
        has_lodge = str(ws.cell(r, 12).value or "").strip() == "是"
        fee = lodging_fee(ws.cell(r, 13).value, ws.cell(r, 14).value)
        tasks.append(
            {
                "task_id": ws.cell(r, 1).value,
                "task_name": ws.cell(r, 2).value,
                "status": status,
                "name": name,
                "biz": biz,
                "station": ws.cell(r, 6).value,
                "pct": pct if pct is not None else 0.0,
                "pct_raw": ws.cell(r, 8).value,
                "start": start,
                "end": end,
                "time_raw": ws.cell(r, 9).value,
                "appeal": appeal,
                "has_lodge": has_lodge or fee > 0,
                "fee": round(fee, 2),
                "apply_fee": ws.cell(r, 13).value,
                "actual_fee": ws.cell(r, 14).value,
                "passed": status in PASS_STATUS,
            }
        )
    return tasks


def load_miles(ws):
    miles = defaultdict(lambda: {"valid": 0.0, "total": 0.0, "invalid": 0.0, "employee_id": "", "biz_user_id": None})
    for r in range(2, ws.max_row + 1):
        name = norm_name(str(ws.cell(r, 3).value or "").strip())
        day = parse_day(ws.cell(r, 7).value)
        if not name or not day:
            continue
        try:
            valid = float(ws.cell(r, 9).value or 0)
        except Exception:
            valid = 0.0
        try:
            total = float(ws.cell(r, 8).value or 0)
        except Exception:
            total = 0.0
        try:
            invalid = float(ws.cell(r, 10).value or 0)
        except Exception:
            invalid = 0.0
        key = (name, day)
        miles[key]["valid"] += valid
        miles[key]["total"] += total
        miles[key]["invalid"] += invalid
        miles[key]["employee_id"] = ws.cell(r, 4).value or miles[key]["employee_id"]
        miles[key]["biz_user_id"] = ws.cell(r, 2).value or miles[key]["biz_user_id"]
    return miles


def person_biz_map(tasks):
    m = defaultdict(set)
    for t in tasks:
        if t["biz"]:
            m[t["name"]].add(t["biz"])
    return m


def resolve_biz(name, day, covering_all, biz_map):
    # 优先当天覆盖任务的业务线；否则用历史任务集合
    if covering_all:
        bizs = {t["biz"] for t in covering_all if t["biz"]}
        if len(bizs) == 1:
            return next(iter(bizs))
        if "大宗" in bizs and "鸡窝" in bizs:
            # 同日双线：取完成度最高那条的业务线
            best = max(covering_all, key=lambda t: (t["pct"], 1 if t["biz"] == "大宗" else 0))
            return best["biz"] or "未知"
        if bizs:
            return next(iter(bizs))
    s = biz_map.get(name, set())
    if s == {"鸡窝"}:
        return "鸡窝"
    if s == {"大宗"}:
        return "大宗"
    if "大宗" in s and "鸡窝" not in s:
        return "大宗"
    if "鸡窝" in s and "大宗" not in s:
        return "鸡窝"
    if "大宗" in s and "鸡窝" in s:
        # 双线人员按日期：8.24前只可能算大宗窗口
        if day < JIWO_START:
            return "大宗"
        return "鸡窝"  # 默认按鸡窝窗口过滤；覆盖任务会再纠正
    return "未知"


def analyze_miles(tasks, miles):
    by_name = defaultdict(list)
    for t in tasks:
        by_name[t["name"]].append(t)
    biz_map = person_biz_map(tasks)

    details = []
    for (name, day), m in sorted(miles.items(), key=lambda x: (x[0][0], x[0][1])):
        valid = round(m["valid"], 2)
        if valid <= 0:
            continue
        covering_all = [t for t in by_name.get(name, []) if t["start"] <= day <= t["end"]]
        biz = resolve_biz(name, day, covering_all, biz_map)
        biz_for_window = biz if biz != "未知" else "大宗"
        if not in_window(biz_for_window, day):
            continue

        # 申诉免扣
        appeal_hit = [t for t in covering_all if t["appeal"]]
        if appeal_hit:
            details.append(
                {
                    "name": name,
                    "employee_id": m["employee_id"],
                    "day": day,
                    "valid": valid,
                    "deduct": 0.0,
                    "keep": valid,
                    "deduct_flag": "否",
                    "reason": "已申诉免扣",
                    "biz": biz,
                    "task_id": appeal_hit[0]["task_id"],
                    "task_time": appeal_hit[0]["time_raw"],
                    "pct_raw": appeal_hit[0]["pct_raw"],
                    "threshold": "",
                    "appeal": "是",
                }
            )
            continue

        covering_pass = [t for t in covering_all if t["passed"]]
        if not covering_pass:
            details.append(
                {
                    "name": name,
                    "employee_id": m["employee_id"],
                    "day": day,
                    "valid": valid,
                    "deduct": valid,
                    "keep": 0.0,
                    "deduct_flag": "是",
                    "reason": "有有效里程但无对应审核通过任务",
                    "biz": biz,
                    "task_id": "",
                    "task_time": "",
                    "pct_raw": "",
                    "threshold": "",
                    "appeal": "否",
                }
            )
            continue

        best = max(covering_pass, key=lambda t: (t["pct"], 1 if t["biz"] == "大宗" else 0))
        thr = THRESH.get(best["biz"], 1.0)
        thr_label = f"{int(thr * 100)}%"
        if best["pct"] + 1e-9 < thr:
            details.append(
                {
                    "name": name,
                    "employee_id": m["employee_id"],
                    "day": day,
                    "valid": valid,
                    "deduct": valid,
                    "keep": 0.0,
                    "deduct_flag": "是",
                    "reason": f"{best['biz']}任务完成度未达标（最高{best['pct_raw']}，阈值{thr_label}）",
                    "biz": best["biz"] or biz,
                    "task_id": best["task_id"],
                    "task_time": best["time_raw"],
                    "pct_raw": best["pct_raw"],
                    "threshold": thr_label,
                    "appeal": "否",
                }
            )
        else:
            details.append(
                {
                    "name": name,
                    "employee_id": m["employee_id"],
                    "day": day,
                    "valid": valid,
                    "deduct": 0.0,
                    "keep": valid,
                    "deduct_flag": "否",
                    "reason": "",
                    "biz": best["biz"] or biz,
                    "task_id": best["task_id"],
                    "task_time": best["time_raw"],
                    "pct_raw": best["pct_raw"],
                    "threshold": thr_label,
                    "appeal": "否",
                }
            )
    return details, by_name, biz_map


def analyze_lodging(tasks):
    rows = []
    for t in tasks:
        if not t["has_lodge"] or t["fee"] <= 0:
            continue
        # 住宿任务需落在业务线窗口内（用任务开始日判断）
        if not in_window(t["biz"] or "大宗", t["start"]):
            continue
        thr = THRESH.get(t["biz"], 1.0)
        fail = (not t["passed"]) or (t["pct"] + 1e-9 < thr)
        if t["appeal"]:
            deduct = 0.0
            reason = "已申诉，住宿费不扣"
        elif fail:
            deduct = t["fee"]
            if not t["passed"]:
                reason = "无审核通过任务/未通过，扣住宿费"
            else:
                reason = f"完成度未达标（{t['pct_raw']}，阈值{int(thr*100)}%），扣住宿费"
        else:
            deduct = 0.0
            reason = "任务达标，住宿费不扣"
        rows.append(
            {
                "name": t["name"],
                "biz": t["biz"],
                "task_id": t["task_id"],
                "task_time": t["time_raw"],
                "status": t["status"],
                "pct_raw": t["pct_raw"],
                "appeal": "是" if t["appeal"] else "否",
                "apply_fee": t["apply_fee"],
                "actual_fee": t["actual_fee"],
                "fee": t["fee"],
                "deduct": deduct,
                "keep": round(t["fee"] - deduct, 2),
                "reason": reason,
            }
        )
    return rows


def summarize(details, lodge_rows, by_name, biz_map):
    mile = defaultdict(lambda: {
        "valid": 0.0, "deduct": 0.0, "keep": 0.0,
        "r1": 0.0, "r2": 0.0, "days_valid": 0, "days_deduct": 0,
        "days_no_task": 0, "days_low_pct": 0, "days_appeal": 0,
        "reasons": set(), "biz": set(), "employee_id": "",
    })
    for d in details:
        s = mile[d["name"]]
        s["valid"] += d["valid"]
        s["deduct"] += d["deduct"]
        s["keep"] += d["keep"]
        s["days_valid"] += 1
        s["employee_id"] = d["employee_id"] or s["employee_id"]
        if d["biz"]:
            s["biz"].add(d["biz"])
        if d["appeal"] == "是":
            s["days_appeal"] += 1
        if d["deduct"] > 0:
            s["days_deduct"] += 1
            s["reasons"].add(d["reason"])
            if "无对应" in d["reason"]:
                s["r1"] += d["deduct"]
                s["days_no_task"] += 1
            else:
                s["r2"] += d["deduct"]
                s["days_low_pct"] += 1

    lodge = defaultdict(lambda: {"fee": 0.0, "deduct": 0.0, "keep": 0.0, "cnt": 0, "deduct_cnt": 0})
    for r in lodge_rows:
        s = lodge[r["name"]]
        s["fee"] += r["fee"]
        s["deduct"] += r["deduct"]
        s["keep"] += r["keep"]
        s["cnt"] += 1
        if r["deduct"] > 0:
            s["deduct_cnt"] += 1

    names = sorted(set(mile) | set(lodge))
    rows = []
    for name in names:
        m = mile[name]
        l = lodge[name]
        biz = "、".join(sorted(m["biz"])) if m["biz"] else ""
        if not biz:
            biz = "、".join(sorted(biz_map.get(name, []))) or "未知"
        rows.append(
            {
                "name": name,
                "employee_id": m["employee_id"],
                "biz": biz,
                "valid": round(m["valid"], 2),
                "mile_deduct": round(m["deduct"], 2),
                "mile_keep": round(m["keep"], 2),
                "r1": round(m["r1"], 2),
                "r2": round(m["r2"], 2),
                "days_valid": m["days_valid"],
                "days_deduct": m["days_deduct"],
                "days_no_task": m["days_no_task"],
                "days_low_pct": m["days_low_pct"],
                "days_appeal": m["days_appeal"],
                "lodge_fee": round(l["fee"], 2),
                "lodge_deduct": round(l["deduct"], 2),
                "lodge_keep": round(l["keep"], 2),
                "reason_summary": "；".join(sorted(m["reasons"])) if m["reasons"] else "",
                "flag_no_task": m["days_no_task"] > 0,
                "flag_low_pct": m["days_low_pct"] > 0,
                "flag_appeal": m["days_appeal"] > 0 or any(
                    t["appeal"] and in_window(t["biz"] or "大宗", t["start"]) for t in by_name.get(name, [])
                ),
            }
        )
    return rows


def stats_block(summary, details, by_name, biz_line):
    """分母：期内有有效里程的人（业务线匹配）。"""
    people = [r for r in summary if r["days_valid"] > 0 and biz_line in (r["biz"] or "")]
    # 双线人员：若 biz 字段含该线，或明细里有该线天数
    if biz_line:
        names_from_detail = {d["name"] for d in details if d["biz"] == biz_line}
        people_names = {r["name"] for r in summary if r["name"] in names_from_detail and r["days_valid"] > 0}
    else:
        people_names = set()

    # rebuild people list strictly by detail biz
    people_names = {d["name"] for d in details if d["biz"] == biz_line}
    n = len(people_names)

    no_task = {d["name"] for d in details if d["biz"] == biz_line and "无对应" in d["reason"]}
    low_pct = {d["name"] for d in details if d["biz"] == biz_line and "完成度未达标" in d["reason"]}
    # 已申诉：窗口内有已申诉任务，或明细免扣
    appeal = set()
    for name in people_names:
        for t in by_name.get(name, []):
            if t["biz"] == biz_line and t["appeal"] and in_window(biz_line, t["start"]):
                appeal.add(name)
                break
        if name not in appeal:
            if any(d["name"] == name and d["biz"] == biz_line and d["appeal"] == "是" for d in details):
                appeal.add(name)

    return {
        "biz": biz_line,
        "people_with_miles": n,
        "no_task_people": len(no_task),
        "low_pct_people": len(low_pct),
        "appeal_people": len(appeal),
        "no_task_names": sorted(no_task),
        "low_pct_names": sorted(low_pct),
        "appeal_names": sorted(appeal),
    }


def style_header(ws, ncol):
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    for c in range(1, ncol + 1):
        cell = ws.cell(1, c)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_out(summary, details, lodge_rows, stats):
    wb = openpyxl.Workbook()
    ws0 = wb.active
    ws0.title = "口径说明"
    for row in [
        ["油补里程 × 任务完成度 / 住宿费扣减分析 v2"],
        ["源文件", str(SRC)],
        ["生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
        [],
        ["规则", "说明"],
        ["里程-无任务", "有有效里程且无「创建/修改审核通过」任务覆盖 → 扣当天全部有效里程"],
        ["里程-完成度", "覆盖任务取最高完成度；大宗<100% 或 鸡窝<75% → 扣当天全部有效里程"],
        ["里程-申诉", "覆盖该天任一条「已申诉」任务 → 当天里程不扣"],
        ["日期-鸡窝", "仅统计 expense_time / 任务开始日 >= 2026-08-24"],
        ["日期-大宗", "统计 2026-08-01 ~ 2026-09-30"],
        ["住宿费金额", "实际住宿费>0用实际，否则用申请费用"],
        ["住宿费扣减", "未申诉且（无审核通过或完成度未达标）→ 扣；已申诉不扣；无有效任务但有住宿费也扣"],
        ["统计分母", "分析窗口内有有效里程的人"],
        ["没创建任务人数", "至少有1天有效里程无审核通过任务覆盖"],
    ]:
        ws0.append(row)
    ws0.column_dimensions["A"].width = 16
    ws0.column_dimensions["B"].width = 88

    # 人数统计
    wsS = wb.create_sheet("人数统计")
    wsS.append(["业务线", "统计窗口", "有有效里程人数", "没创建任务人数", "完成度未达标人数", "已申诉人数"])
    for st, window in [
        (stats["鸡窝"], "2026-08-24起"),
        (stats["大宗"], "2026-08 ~ 2026-09"),
    ]:
        wsS.append([st["biz"], window, st["people_with_miles"], st["no_task_people"], st["low_pct_people"], st["appeal_people"]])
    style_header(wsS, 6)
    for i, w in enumerate([10, 16, 16, 16, 16, 12], 1):
        wsS.column_dimensions[get_column_letter(i)].width = w
    wsS.append([])
    wsS.append(["附：名单明细（逗号分隔）"])
    for st in [stats["鸡窝"], stats["大宗"]]:
        wsS.append([st["biz"], "没创建任务", "、".join(st["no_task_names"])])
        wsS.append([st["biz"], "完成度未达标", "、".join(st["low_pct_names"])])
        wsS.append([st["biz"], "已申诉", "、".join(st["appeal_names"])])

    # 按人汇总
    ws1 = wb.create_sheet("按人汇总")
    h1 = [
        "业务员", "工号", "业务线",
        "有效里程合计", "扣减里程", "可报销里程", "其中无任务扣减", "其中完成度未达标扣减",
        "有里程天数", "扣减天数", "无任务天数", "完成度未达标天数", "申诉免扣天数",
        "住宿费合计", "扣减住宿费", "可报住宿费",
        "是否有无任务天", "是否有完成度未达标", "是否已申诉",
        "里程扣减原因汇总",
    ]
    ws1.append(h1)
    for r in summary:
        ws1.append([
            r["name"], r["employee_id"], r["biz"],
            r["valid"], r["mile_deduct"], r["mile_keep"], r["r1"], r["r2"],
            r["days_valid"], r["days_deduct"], r["days_no_task"], r["days_low_pct"], r["days_appeal"],
            r["lodge_fee"], r["lodge_deduct"], r["lodge_keep"],
            "是" if r["flag_no_task"] else "否",
            "是" if r["flag_low_pct"] else "否",
            "是" if r["flag_appeal"] else "否",
            r["reason_summary"],
        ])
    style_header(ws1, len(h1))
    for i, w in enumerate([10, 10, 10, 12, 10, 10, 12, 14, 10, 8, 10, 12, 10, 10, 10, 10, 12, 14, 10, 50], 1):
        ws1.column_dimensions[get_column_letter(i)].width = w
    if summary:
        ws1.append([
            "合计", "", "",
            round(sum(x["valid"] for x in summary), 2),
            round(sum(x["mile_deduct"] for x in summary), 2),
            round(sum(x["mile_keep"] for x in summary), 2),
            round(sum(x["r1"] for x in summary), 2),
            round(sum(x["r2"] for x in summary), 2),
            sum(x["days_valid"] for x in summary),
            sum(x["days_deduct"] for x in summary),
            sum(x["days_no_task"] for x in summary),
            sum(x["days_low_pct"] for x in summary),
            sum(x["days_appeal"] for x in summary),
            round(sum(x["lodge_fee"] for x in summary), 2),
            round(sum(x["lodge_deduct"] for x in summary), 2),
            round(sum(x["lodge_keep"] for x in summary), 2),
            "", "", "", "",
        ])
        for c in range(1, len(h1) + 1):
            ws1.cell(ws1.max_row, c).font = Font(bold=True)

    # 按天明细
    ws2 = wb.create_sheet("按天明细")
    h2 = ["业务员", "工号", "日期", "业务线", "有效里程", "是否扣减", "扣减里程", "可报销里程",
          "扣减原因", "是否申诉免扣", "关联任务编号", "任务时间", "任务完成度", "达标阈值"]
    ws2.append(h2)
    for d in details:
        ws2.append([
            d["name"], d["employee_id"], d["day"].isoformat(), d["biz"], d["valid"],
            d["deduct_flag"], d["deduct"], d["keep"], d["reason"], d["appeal"],
            d["task_id"], d["task_time"], d["pct_raw"], d["threshold"],
        ])
    style_header(ws2, len(h2))
    for i, w in enumerate([10, 10, 12, 8, 10, 8, 10, 10, 42, 10, 14, 22, 10, 8], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    ws3 = wb.create_sheet("仅扣减明细")
    ws3.append(h2)
    for d in details:
        if d["deduct"] > 0:
            ws3.append([
                d["name"], d["employee_id"], d["day"].isoformat(), d["biz"], d["valid"],
                d["deduct_flag"], d["deduct"], d["keep"], d["reason"], d["appeal"],
                d["task_id"], d["task_time"], d["pct_raw"], d["threshold"],
            ])
    style_header(ws3, len(h2))
    for i, w in enumerate([10, 10, 12, 8, 10, 8, 10, 10, 42, 10, 14, 22, 10, 8], 1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    # 住宿费明细
    ws4 = wb.create_sheet("住宿费明细")
    h4 = ["业务员", "业务线", "任务编号", "任务时间", "审核状态", "完成度", "是否申诉",
          "申请费用", "实际住宿费", "计费金额", "扣减住宿费", "可报住宿费", "说明"]
    ws4.append(h4)
    for r in lodge_rows:
        ws4.append([
            r["name"], r["biz"], r["task_id"], r["task_time"], r["status"], r["pct_raw"], r["appeal"],
            r["apply_fee"], r["actual_fee"], r["fee"], r["deduct"], r["keep"], r["reason"],
        ])
    style_header(ws4, len(h4))
    for i, w in enumerate([10, 8, 14, 22, 12, 8, 8, 10, 10, 10, 10, 10, 36], 1):
        ws4.column_dimensions[get_column_letter(i)].width = w

    wb.save(OUT)
    try:
        wb.save(OUT2)
    except Exception as e:
        print("backup save fail", e)
    try:
        wb.save(OUT3)
    except Exception as e:
        print("desktop overwrite fail (file may be open):", e)
    return OUT


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    tasks = load_tasks(wb["8月业务员创建任务和完成情况"])
    miles = load_miles(wb["拜访里程"])
    details, by_name, biz_map = analyze_miles(tasks, miles)
    lodge_rows = analyze_lodging(tasks)
    summary = summarize(details, lodge_rows, by_name, biz_map)
    stats = {
        "鸡窝": stats_block(summary, details, by_name, "鸡窝"),
        "大宗": stats_block(summary, details, by_name, "大宗"),
    }
    path = write_out(summary, details, lodge_rows, stats)

    print("details", len(details), "deduct_days", sum(1 for d in details if d["deduct"] > 0))
    print("mile_valid", round(sum(d["valid"] for d in details), 2))
    print("mile_deduct", round(sum(d["deduct"] for d in details), 2))
    print("mile_keep", round(sum(d["keep"] for d in details), 2))
    print("lodge_fee", round(sum(r["fee"] for r in lodge_rows), 2))
    print("lodge_deduct", round(sum(r["deduct"] for r in lodge_rows), 2))
    for k, st in stats.items():
        print(k, "people", st["people_with_miles"], "no_task", st["no_task_people"],
              "low_pct", st["low_pct_people"], "appeal", st["appeal_people"])
    print("saved", path)
    # 张桐菓 / 王志刚 核对
    for who in ("张桐菓", "王志刚"):
        zt = [r for r in summary if r["name"] == who]
        print(who, "汇总", zt)
        zt_d = [d for d in details if d["name"] == who]
        print(who, "人天", len(zt_d), "扣减天", sum(1 for d in zt_d if d["deduct"] > 0))
        for d in zt_d:
            print(d["day"], "valid", d["valid"], "deduct", d["deduct"], d["reason"] or "保留", d["task_id"], d["pct_raw"])
        zt_l = [r for r in lodge_rows if r["name"] == who]
        print(who, "住宿", zt_l)


if __name__ == "__main__":
    main()
