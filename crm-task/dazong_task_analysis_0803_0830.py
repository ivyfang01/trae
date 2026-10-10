# -*- coding: utf-8 -*-
"""大宗专项分析：2026-08-03 ~ 2026-08-30（对齐鸡窝报告模板）"""
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime, date
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

SRC = Path(r"c:\Users\HY\Desktop\cursor分析数据原始表.xlsx")
OUT = Path(r"c:\Users\HY\Desktop\大宗任务分析_0803_0830.xlsx")
OUT2 = Path(r"d:\trae\crm-task\大宗任务分析_0803_0830.xlsx")

PASS = {"创建审核通过", "修改审核通过"}
CREATE_STATUS = {"创建审核通过", "创建审核中", "创建审核驳回"}
MODIFY_STATUS = {"修改审核通过", "修改审核中"}
START = date(2026, 8, 3)
END = date(2026, 8, 30)
THRESH = 1.0
BIZ = "大宗"
NAME_ALIASES = {"张桐?": "张桐菓", "王志刚1": "王志刚"}


def norm(n):
    return NAME_ALIASES.get((n or "").strip(), (n or "").strip())


def parse_pct(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        x = float(v)
        return x / 100.0 if x > 1 else x
    try:
        return float(str(v).strip().replace("%", "")) / 100.0
    except Exception:
        return None


def parse_day(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()[:10]
    for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except Exception:
            continue
    return None


def parse_range(s):
    if s is None:
        return None, None
    text = str(s).strip()
    if " ~ " in text:
        a, b = text.split(" ~ ", 1)
    elif "~" in text:
        a, b = text.split("~", 1)
    else:
        d = parse_day(text)
        return d, d
    return parse_day(a.strip()), parse_day(b.strip())


def overlap(a1, a2, b1, b2):
    return a1 <= b2 and b1 <= a2


def lodging_fee(apply_fee, actual_fee):
    try:
        actual = float(actual_fee or 0)
    except Exception:
        actual = 0.0
    try:
        apply = float(apply_fee or 0)
    except Exception:
        apply = 0.0
    return actual if actual > 0 else apply


def load(wb):
    ws = wb["8月业务员创建任务和完成情况"]
    tasks = []
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, 5).value or "").strip() != BIZ:
            continue
        name = norm(str(ws.cell(r, 4).value or ""))
        start, end = parse_range(ws.cell(r, 9).value)
        if not name or not start or not end:
            continue
        if end < start:
            start, end = end, start
        if not overlap(start, end, START, END):
            continue
        status = str(ws.cell(r, 3).value or "").strip()
        pct = parse_pct(ws.cell(r, 8).value)
        fee = lodging_fee(ws.cell(r, 13).value, ws.cell(r, 14).value)
        tasks.append(
            {
                "task_id": ws.cell(r, 1).value,
                "name": name,
                "status": status,
                "pct": pct if pct is not None else 0.0,
                "pct_raw": ws.cell(r, 8).value,
                "start": start,
                "end": end,
                "time_raw": ws.cell(r, 9).value,
                "appeal": str(ws.cell(r, 10).value or "").strip() == "已申诉",
                "passed": status in PASS,
                "is_create": status in CREATE_STATUS,
                "is_modify": status in MODIFY_STATUS,
                "fee": round(fee, 2),
                "has_lodge": str(ws.cell(r, 12).value or "").strip() == "是" or fee > 0,
                "apply_fee": ws.cell(r, 13).value,
                "actual_fee": ws.cell(r, 14).value,
                "station": ws.cell(r, 6).value,
            }
        )

    all_dazong = set()
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, 5).value or "").strip() == BIZ:
            all_dazong.add(norm(str(ws.cell(r, 4).value or "")))

    ws2 = wb["拜访里程"]
    miles = defaultdict(lambda: {"valid": 0.0, "employee_id": ""})
    for r in range(2, ws2.max_row + 1):
        name = norm(str(ws2.cell(r, 3).value or ""))
        day = parse_day(ws2.cell(r, 7).value)
        if not name or not day or not (START <= day <= END):
            continue
        try:
            valid = float(ws2.cell(r, 9).value or 0)
        except Exception:
            valid = 0.0
        if valid <= 0:
            continue
        miles[(name, day)]["valid"] += valid
        miles[(name, day)]["employee_id"] = ws2.cell(r, 4).value or miles[(name, day)]["employee_id"]

    return tasks, miles, all_dazong


def analyze(tasks, miles, all_dazong):
    by_name = defaultdict(list)
    for t in tasks:
        by_name[t["name"]].append(t)

    mile_people = sorted({n for (n, d) in miles})
    universe = sorted({n for n in mile_people if n in all_dazong or n in by_name})

    has_any_task = set(by_name.keys())
    created_only = {t["name"] for t in tasks if t["is_create"]}
    modified = {t["name"] for t in tasks if t["is_modify"]}
    created_people = has_any_task
    no_create = [n for n in universe if n not in has_any_task]

    below = set()
    for name, ts in by_name.items():
        for t in ts:
            if t["passed"] and t["pct"] + 1e-9 < THRESH:
                below.add(name)
                break

    appealed = {t["name"] for t in tasks if t["appeal"]}

    details = []
    for (name, day), m in sorted(miles.items()):
        if name not in all_dazong and name not in by_name:
            continue
        valid = round(m["valid"], 2)
        covering_all = [t for t in by_name.get(name, []) if t["start"] <= day <= t["end"]]
        if any(t["appeal"] for t in covering_all):
            details.append(
                {
                    "name": name,
                    "employee_id": m["employee_id"],
                    "day": day,
                    "valid": valid,
                    "deduct": 0.0,
                    "reason": "已申诉免扣",
                    "appeal": "是",
                }
            )
            continue
        covering_pass = [t for t in covering_all if t["passed"]]
        if not covering_pass:
            details.append(
                {
                    "name": name,
                    "employee_id": m["employee_id"],
                    "day": day,
                    "valid": valid,
                    "deduct": valid,
                    "reason": "有有效里程但无对应审核通过任务",
                    "appeal": "否",
                }
            )
            continue
        best = max(covering_pass, key=lambda t: t["pct"])
        if best["pct"] + 1e-9 < THRESH:
            details.append(
                {
                    "name": name,
                    "employee_id": m["employee_id"],
                    "day": day,
                    "valid": valid,
                    "deduct": valid,
                    "reason": f"完成度未达标（最高{best['pct_raw']}，阈值100%）",
                    "appeal": "否",
                    "task_id": best["task_id"],
                    "pct_raw": best["pct_raw"],
                }
            )
        else:
            details.append(
                {
                    "name": name,
                    "employee_id": m["employee_id"],
                    "day": day,
                    "valid": valid,
                    "deduct": 0.0,
                    "reason": "",
                    "appeal": "否",
                    "task_id": best["task_id"],
                    "pct_raw": best["pct_raw"],
                }
            )

    below_by_day = {d["name"] for d in details if "完成度未达标" in d["reason"]}
    below_people = sorted(below | below_by_day)

    lodge_rows = []
    for t in tasks:
        if not t["has_lodge"] or t["fee"] <= 0:
            continue
        fail = (not t["passed"]) or (t["pct"] + 1e-9 < THRESH)
        if t["appeal"]:
            deduct = 0.0
            reason = "已申诉，住宿费不扣"
        elif fail:
            deduct = t["fee"]
            reason = "未通过或完成度未达标，扣住宿费" if t["passed"] else "无审核通过任务，扣住宿费"
        else:
            deduct = 0.0
            reason = "任务达标，住宿费不扣"
        lodge_rows.append({**t, "deduct": deduct, "keep": round(t["fee"] - deduct, 2), "reason": reason})

    # 任务状态 / 完成度分布（报告用）
    status_cnt = Counter(t["status"] for t in tasks)
    pass_pct_cnt = Counter()
    for t in tasks:
        if t["passed"]:
            raw = t["pct_raw"]
            if isinstance(raw, (int, float)):
                label = f"{int(round(raw * 100)) if raw <= 1 else int(round(raw))}%"
            else:
                label = str(raw).strip() if raw not in (None, "") else f"{int(round(t['pct'] * 100))}%"
            if not label.endswith("%"):
                label = f"{label}%"
            pass_pct_cnt[label] += 1

    no_pass_km = round(sum(d["deduct"] for d in details if "无对应审核通过" in d["reason"]), 2)
    low_pct_km = round(sum(d["deduct"] for d in details if "完成度未达标" in d["reason"]), 2)
    no_pass_days = sum(1 for d in details if "无对应审核通过" in d["reason"])
    low_pct_days = sum(1 for d in details if "完成度未达标" in d["reason"])

    person_deduct = defaultdict(float)
    for d in details:
        if d["deduct"] > 0:
            person_deduct[d["name"]] += d["deduct"]
    top10 = sorted(person_deduct.items(), key=lambda x: -x[1])[:10]
    people_with_deduct = len(person_deduct)

    summary = {
        "window": f"{START} ~ {END}",
        "universe": len(universe),
        "created": len(created_people),
        "no_create": len(no_create),
        "modified": len(modified),
        "below100": len(below_people),
        "appealed": len(appealed),
        "mile_deduct": round(sum(d["deduct"] for d in details), 2),
        "mile_valid": round(sum(d["valid"] for d in details), 2),
        "lodge_deduct": round(sum(r["deduct"] for r in lodge_rows), 2),
        "lodge_fee": round(sum(r["fee"] for r in lodge_rows), 2),
        "task_total": len(tasks),
        "status_cnt": dict(status_cnt),
        "pass_pct_cnt": dict(pass_pct_cnt.most_common()),
        "no_pass_km": no_pass_km,
        "low_pct_km": low_pct_km,
        "no_pass_days": no_pass_days,
        "low_pct_days": low_pct_days,
        "top10": top10,
        "people_with_deduct": people_with_deduct,
        "created_names": sorted(created_people),
        "no_create_names": no_create,
        "modified_names": sorted(modified),
        "below_names": below_people,
        "appeal_names": sorted(appealed),
        "created_status_create": sorted(created_only),
    }
    return summary, details, lodge_rows, universe, by_name


def style_header(ws, n):
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    for c in range(1, n + 1):
        cell = ws.cell(1, c)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)


def write(summary, details, lodge_rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "大宗汇总"
    rows = [
        ["大宗任务专项分析"],
        ["统计窗口", summary["window"]],
        ["生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
        [],
        ["指标", "人数/金额", "口径说明"],
        ["有有效里程的大宗人数（分母）", summary["universe"], "8.3-8.30 有有效里程，且任务表业务线为大宗"],
        ["创建了任务", summary["created"], "窗口内有任意大宗任务（含创建/修改各类状态）"],
        ["没创建任务", summary["no_create"], "分母中窗口内一条大宗任务都没有"],
        ["修改了任务", summary["modified"], "窗口内存在「修改审核通过/修改审核中」任务"],
        ["未达标100%", summary["below100"], "存在审核通过且完成度<100%的任务，或按天判定未达标"],
        ["已申诉", summary["appealed"], "窗口内任务「是否申诉=已申诉」"],
        ["需扣减里程", summary["mile_deduct"], "无通过任务 / 完成度<100%；已申诉免扣"],
        ["窗口有效里程合计", summary["mile_valid"], "对照用"],
        ["需扣减住宿费", summary["lodge_deduct"], "未申诉且（未通过或完成度<100%）；金额=实际优先否则申请"],
        ["窗口住宿费计费合计", summary["lodge_fee"], "对照用"],
        ["窗口任务条数", summary["task_total"], "任务时间与窗口有交集的大宗任务"],
    ]
    for r in rows:
        ws.append(r)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 56

    def list_sheet(title, names, header="业务员"):
        w = wb.create_sheet(title)
        w.append([header])
        style_header(w, 1)
        for n in names:
            w.append([n])
        w.column_dimensions["A"].width = 16
        w.append([])
        w.append(["人数", len(names)])

    list_sheet("名单_创建了任务", summary["created_names"])
    list_sheet("名单_没创建", summary["no_create_names"])
    list_sheet("名单_修改了任务", summary["modified_names"])
    list_sheet("名单_未达标100%", summary["below_names"])
    list_sheet("名单_已申诉", summary["appeal_names"])

    w2 = wb.create_sheet("里程扣减明细")
    h = ["业务员", "工号", "日期", "有效里程", "扣减里程", "原因", "是否申诉免扣"]
    w2.append(h)
    style_header(w2, len(h))
    for d in details:
        if d["deduct"] <= 0:
            continue
        w2.append(
            [
                d["name"],
                d.get("employee_id"),
                d["day"].isoformat(),
                d["valid"],
                d["deduct"],
                d["reason"],
                d["appeal"],
            ]
        )
    for i, width in enumerate([12, 12, 12, 10, 10, 42, 12], 1):
        w2.column_dimensions[get_column_letter(i)].width = width

    w3 = wb.create_sheet("住宿费扣减明细")
    h3 = ["业务员", "任务编号", "任务时间", "审核状态", "完成度", "是否申诉", "计费金额", "扣减住宿费", "说明"]
    w3.append(h3)
    style_header(w3, len(h3))
    for r in lodge_rows:
        if r["deduct"] <= 0:
            continue
        w3.append(
            [
                r["name"],
                r["task_id"],
                r["time_raw"],
                r["status"],
                r["pct_raw"],
                "是" if r["appeal"] else "否",
                r["fee"],
                r["deduct"],
                r["reason"],
            ]
        )
    for i, width in enumerate([12, 14, 24, 14, 8, 8, 10, 10, 36], 1):
        w3.column_dimensions[get_column_letter(i)].width = width

    w4 = wb.create_sheet("任务状态分布")
    w4.append(["审核状态", "条数"])
    style_header(w4, 2)
    for k, v in sorted(summary["status_cnt"].items(), key=lambda x: -x[1]):
        w4.append([k, v])

    w5 = wb.create_sheet("通过任务完成度分布")
    w5.append(["完成度", "条数"])
    style_header(w5, 2)
    for k, v in summary["pass_pct_cnt"].items():
        w5.append([k, v])

    wb.save(OUT)
    try:
        wb.save(OUT2)
    except Exception as e:
        print("backup fail", e)
    return OUT


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    tasks, miles, all_dazong = load(wb)
    summary, details, lodge_rows, universe, by_name = analyze(tasks, miles, all_dazong)
    path = write(summary, details, lodge_rows)

    # 打印报告所需关键数字
    import json

    dump = {
        k: v
        for k, v in summary.items()
        if k
        not in (
            "created_names",
            "no_create_names",
            "modified_names",
            "below_names",
            "appeal_names",
            "created_status_create",
        )
    }
    dump["created_names"] = summary["created_names"]
    dump["no_create_names"] = summary["no_create_names"]
    dump["modified_names"] = summary["modified_names"]
    dump["below_names"] = summary["below_names"]
    dump["appeal_names"] = summary["appeal_names"]
    dump["top10"] = [(n, round(km, 2)) for n, km in summary["top10"]]
    print(json.dumps(dump, ensure_ascii=False, indent=2, default=str))
    print("saved", path)


if __name__ == "__main__":
    main()
