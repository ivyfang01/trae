# -*- coding: utf-8 -*-
"""填充「是否有对应的任务」：里程 expense_time 是否落在该人任一任务时间区间内。"""
from __future__ import annotations

from collections import defaultdict
from datetime import datetime, date
from pathlib import Path

import openpyxl

MILE_PATH = Path(r"c:\Users\HY\Desktop\8月业务员里程数.xlsx")
TASK_PATH = Path(r"c:\Users\HY\Desktop\8月任务列数据.xlsx")
OUT_BACKUP = Path(r"d:\trae\crm-task\8月业务员里程数_已填任务对应.xlsx")

NAME_ALIASES = {"张桐?": "张桐菓", "王志刚1": "王志刚"}
PASS_ALL = True  # 不区分审核状态
JIWO_START = date(2026, 8, 24)
JIWO_END = date(2026, 8, 31)
DZ_START = date(2026, 8, 1)
DZ_END = date(2026, 8, 31)


def norm(n: str) -> str:
    n = (n or "").strip()
    return NAME_ALIASES.get(n, n)


def parse_day(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()[:10]
    for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except Exception:
            continue
    return None


def parse_range(s):
    if s is None:
        return None, None
    text = str(s).strip()
    if " ~ " in text:
        a, b = text.split(" ~ ", 1)
    elif "~" in text:
        a, b = text.split("~", 1)
    else:
        d = parse_day(text)
        return d, d
    return parse_day(a.strip()), parse_day(b.strip())


def task_kind(name: str) -> str:
    """日度 / 周度 / 其他"""
    name = name or ""
    if "周" in name:
        return "weekly"
    if "日" in name:
        return "daily"
    return "other"


def load_tasks():
    wb = openpyxl.load_workbook(TASK_PATH, data_only=True)
    ws = wb["0"]
    # person -> list of (start, end, kind)
    by_name = defaultdict(list)
    kinds = defaultdict(set)
    n = 0
    for r in ws.iter_rows(min_row=2, values_only=True):
        person = norm(str(r[3] or ""))
        start, end = parse_range(r[6])
        if not person or not start or not end:
            continue
        if end < start:
            start, end = end, start
        kind = task_kind(str(r[1] or ""))
        by_name[person].append((start, end, kind))
        kinds[person].add(kind)
        n += 1
    wb.close()
    return by_name, kinds, n


def person_window(kinds_set: set[str]):
    """返回 (eval_start, eval_end)；无任务则用整月大宗窗。"""
    has_daily = "daily" in kinds_set
    has_weekly = "weekly" in kinds_set
    if has_daily and not has_weekly:
        return JIWO_START, JIWO_END  # 鸡窝
    if has_weekly and not has_daily:
        return DZ_START, DZ_END  # 大宗
    if has_daily and has_weekly:
        return DZ_START, DZ_END  # 双线：整月，任一日度/周度覆盖都算
    return DZ_START, DZ_END  # 无任务：整月评，结果均为否


def covered(ranges, day: date) -> bool:
    for s, e, _k in ranges:
        if s <= day <= e:
            return True
    return False


def main():
    by_name, kinds, n_task = load_tasks()
    print("tasks loaded", n_task, "people", len(by_name))

    wb = openpyxl.load_workbook(MILE_PATH)
    ws = wb[wb.sheetnames[0]]
    # find last col header
    headers = [c.value for c in ws[1]]
    try:
        col = headers.index("是否有对应的任务") + 1
    except ValueError:
        col = len(headers) + 1
        ws.cell(1, col, "是否有对应的任务")

    stats = {"是": 0, "否": 0, "空": 0}
    samples_no = []

    for r in range(2, ws.max_row + 1):
        name = norm(str(ws.cell(r, 8).value or ""))  # biz_user_name
        day = parse_day(ws.cell(r, 13).value)  # expense_time
        if not name or not day:
            ws.cell(r, col, None)
            stats["空"] += 1
            continue

        kset = kinds.get(name, set())
        w0, w1 = person_window(kset)
        if not (w0 <= day <= w1):
            # 鸡窝窗口外：不考核，留空
            ws.cell(r, col, None)
            stats["空"] += 1
            continue

        ranges = by_name.get(name, [])
        yes = covered(ranges, day)
        val = "是" if yes else "否"
        ws.cell(r, col, val)
        stats[val] += 1
        if val == "否" and len(samples_no) < 8:
            samples_no.append((name, day.isoformat(), sorted(kset)))

    try:
        wb.save(MILE_PATH)
        saved = MILE_PATH
    except PermissionError:
        wb.save(OUT_BACKUP)
        saved = OUT_BACKUP
        print("desktop locked, saved backup", OUT_BACKUP)
    else:
        try:
            wb.save(OUT_BACKUP)
        except Exception as e:
            print("backup skip", e)

    print("stats", stats)
    print("sample 否", samples_no)
    print("saved", saved)


if __name__ == "__main__":
    main()
