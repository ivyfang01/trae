# -*- coding: utf-8 -*-
"""按表内「业务类型」重填「是否有对应的任务」。

规则：
- 鸡窝：仅 8.24–8.31；expense_time 落在该人任一日度/周度任务区间 → 是，否则 否；窗口外留空
- 大宗：8.1–8.31；同上
- 采购/技术老师/蛋品/其他：留空
- 不区分审核状态；张桐?↔张桐菓
"""
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime, date
from pathlib import Path

import openpyxl

MILE = Path(r"c:\Users\HY\Desktop\8月业务员里程数_已填任务对应.xlsx")
TASK = Path(r"c:\Users\HY\Desktop\8月任务列数据.xlsx")
OUT_BACKUP = Path(r"d:\trae\crm-task\8月业务员里程数_已填任务对应.xlsx")

NAME_ALIASES = {"张桐?": "张桐菓", "王志刚1": "王志刚"}
JIWO_START, JIWO_END = date(2026, 8, 24), date(2026, 8, 31)
DZ_START, DZ_END = date(2026, 8, 1), date(2026, 8, 31)
EVAL_BIZ = {"鸡窝", "大宗"}


def norm(n: str) -> str:
    n = (n or "").strip()
    return NAME_ALIASES.get(n, n)


def parse_day(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()[:10]
    for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except Exception:
            continue
    return None


def parse_range(s):
    if s is None:
        return None, None
    text = str(s).strip()
    if " ~ " in text:
        a, b = text.split(" ~ ", 1)
    elif "~" in text:
        a, b = text.split("~", 1)
    else:
        d = parse_day(text)
        return d, d
    return parse_day(a.strip()), parse_day(b.strip())


def load_tasks():
    wb = openpyxl.load_workbook(TASK, data_only=True)
    ws = wb["0"]
    by_name = defaultdict(list)
    for r in ws.iter_rows(min_row=2, values_only=True):
        person = norm(str(r[3] or ""))
        start, end = parse_range(r[6])
        if not person or not start or not end:
            continue
        if end < start:
            start, end = end, start
        by_name[person].append((start, end))
    wb.close()
    return by_name


def window_for(biz: str):
    if biz == "鸡窝":
        return JIWO_START, JIWO_END
    if biz == "大宗":
        return DZ_START, DZ_END
    return None


def covered(ranges, day: date) -> bool:
    return any(s <= day <= e for s, e in ranges)


def main():
    by_name = load_tasks()
    wb = openpyxl.load_workbook(MILE)
    ws = wb[wb.sheetnames[0]]
    headers = [c.value for c in ws[1]]
    col_name = headers.index("biz_user_name") + 1
    col_day = headers.index("expense_time") + 1
    col_biz = headers.index("业务类型") + 1
    col_flag = headers.index("是否有对应的任务") + 1

    stats = Counter()
    by_biz = Counter()
    samples = []

    for r in range(2, ws.max_row + 1):
        name = norm(str(ws.cell(r, col_name).value or ""))
        day = parse_day(ws.cell(r, col_day).value)
        biz = str(ws.cell(r, col_biz).value or "").strip()

        if biz not in EVAL_BIZ or not name or not day:
            ws.cell(r, col_flag, None)
            stats["空_非考核"] += 1
            continue

        win = window_for(biz)
        if not (win[0] <= day <= win[1]):
            ws.cell(r, col_flag, None)
            stats["空_窗口外"] += 1
            by_biz[(biz, "空_窗口外")] += 1
            continue

        val = "是" if covered(by_name.get(name, []), day) else "否"
        ws.cell(r, col_flag, val)
        stats[val] += 1
        by_biz[(biz, val)] += 1
        if val == "否" and len(samples) < 6:
            samples.append((name, biz, day.isoformat()))

    saved = []
    for out in (MILE, OUT_BACKUP):
        try:
            wb.save(out)
            saved.append(str(out))
        except PermissionError:
            alt = out.with_name(out.stem + "_重填任务对应.xlsx")
            wb.save(alt)
            saved.append(str(alt) + " (原文件占用)")

    print("stats", dict(stats))
    print("by_biz", {f"{a}|{b}": n for (a, b), n in sorted(by_biz.items())})
    print("sample 否", samples)
    print("saved", saved)


if __name__ == "__main__":
    main()
