# -*- coding: utf-8 -*-
"""Generate CRM任务系统闭环 PPT matching reference slide."""
from io import BytesIO
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from lxml import etree

OUT = Path(r"d:\trae\crm-task\CRM任务系统闭环.pptx")
REF = Path(
    r"C:\Users\HY\.cursor\projects\d\assets\c__Users_HY_AppData_Roaming_Cursor_User_workspaceStorage_4329c1b5e94952ef11948a41effc4ef4_images_image-58b56aba-4047-494f-9d54-0c0a881de602.png"
)

COLORS = {
    "01": RGBColor(0x6A, 0xAB, 0xE8),
    "02": RGBColor(0xF0, 0x8C, 0x45),
    "03": RGBColor(0xF5, 0xC8, 0x42),
    "04": RGBColor(0x8B, 0xC3, 0x4A),
    "05": RGBColor(0x4E, 0xCD, 0xC4),
    "06": RGBColor(0xF0, 0x78, 0x78),
}

STEPS = [
    {"id": "01", "title": "01 发布任务", "desc": "明确工作计划,待办与日程同步", "side": "left", "y": 1.55},
    {"id": "06", "title": "06 数据看板", "desc": "任务、出勤、绩效、拜访转化数据", "side": "left", "y": 3.35},
    {"id": "05", "title": "05 日报复盘", "desc": "当日工作汇总,持续改善", "side": "left", "y": 5.15},
    {"id": "02", "title": "02 客户拜访", "desc": "客户标签、素材库、陪访、里程申诉", "side": "right", "y": 1.55},
    {"id": "03", "title": "03 更新客户档案", "desc": "档案信息线上化,完善客户画像", "side": "right", "y": 3.35},
    {"id": "04", "title": "04 机会点跟进", "desc": "挖掘潜在价值,精准推送服务", "side": "right", "y": 5.15},
]


def crop_center_ring(ref_path: Path) -> bytes:
    """Crop center diagram from reference image for pixel-accurate ring."""
    im = Image.open(ref_path).convert("RGBA")
    w, h = im.size
    # center square crop covering the ring
    size = int(min(w, h) * 0.52)
    left = (w - size) // 2
    top = int(h * 0.18)
    ring = im.crop((left, top, left + size, top + size))
    buf = BytesIO()
    ring.save(buf, format="PNG")
    return buf.getvalue()


def set_dashed_line(connector, color: RGBColor):
    sp_pr = connector._element.spPr
    ln_el = sp_pr.find(qn("a:ln"))
    if ln_el is None:
        ln_el = etree.SubElement(sp_pr, qn("a:ln"))
    prst = ln_el.find(qn("a:prstDash"))
    if prst is None:
        prst = etree.SubElement(ln_el, qn("a:prstDash"))
    prst.set("val", "sysDot")
    solid = ln_el.find(qn("a:solidFill"))
    if solid is None:
        solid = etree.SubElement(ln_el, qn("a:solidFill"))
    srgb = solid.find(qn("a:srgbClr"))
    if srgb is None:
        srgb = etree.SubElement(solid, qn("a:srgbClr"))
    srgb.set("val", f"{color[0]:02X}{color[1]:02X}{color[2]:02X}")
    ln_el.set("w", "9525")  # ~0.75pt


def add_text_block(slide, step):
    color_rgb = COLORS[step["id"]]
    if step["side"] == "left":
        left = Inches(0.55)
        align = PP_ALIGN.LEFT
    else:
        left = Inches(9.75)
        align = PP_ALIGN.LEFT

    box = slide.shapes.add_textbox(left, Inches(step["y"]), Inches(3.05), Inches(0.95))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    p1 = tf.paragraphs[0]
    p1.alignment = align
    r1 = p1.add_run()
    r1.text = step["title"]
    r1.font.size = Pt(22)
    r1.font.bold = True
    r1.font.name = "Microsoft YaHei"
    r1.font.color.rgb = color_rgb

    p2 = tf.add_paragraph()
    p2.alignment = align
    p2.space_before = Pt(2)
    r2 = p2.add_run()
    r2.text = step["desc"]
    r2.font.size = Pt(13)
    r2.font.name = "Microsoft YaHei"
    r2.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def add_dot(slide, x, y, color: RGBColor, size=Inches(0.07)):
    shape = slide.shapes.add_shape(1, int(x - size / 2), int(y - size / 2), int(size), int(size))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def build_ppt():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.55), Inches(0.38), Inches(7), Inches(0.65))
    tr = title_box.text_frame.paragraphs[0].add_run()
    tr.text = "CRM任务系统闭环"
    tr.font.size = Pt(34)
    tr.font.bold = True
    tr.font.name = "SimSun"
    tr.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    # Center ring from reference crop
    ring_png = crop_center_ring(REF)
    ring_size = Inches(4.55)
    ring_left = (prs.slide_width - ring_size) / 2
    ring_top = Inches(1.48)
    slide.shapes.add_picture(BytesIO(ring_png), ring_left, ring_top, width=ring_size, height=ring_size)

    center_x = prs.slide_width / 2

    for step in STEPS:
        add_text_block(slide, step)
        color_rgb = COLORS[step["id"]]

        if step["side"] == "left":
            dot_x = Inches(3.72)
            end_x = center_x - Inches(1.02)
            begin_x = dot_x + Inches(0.04)
        else:
            dot_x = Inches(9.58)
            end_x = center_x + Inches(1.02)
            begin_x = dot_x - Inches(0.04)

        dot_y = Inches(step["y"] + 0.2)
        add_dot(slide, dot_x, dot_y, color_rgb)

        conn = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            int(begin_x),
            int(dot_y + Inches(0.035)),
            int(end_x),
            int(dot_y + Inches(0.035)),
        )
        set_dashed_line(conn, color_rgb)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"saved: {OUT}")


if __name__ == "__main__":
    build_ppt()
