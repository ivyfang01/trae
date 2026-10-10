# -*- coding: utf-8 -*-
"""为里程表增加业务类型列：鸡窝/大宗/采购/技术老师"""
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime, date
from pathlib import Path

import openpyxl
import pymysql
from pymysql.cursors import DictCursor

MILE = Path(r"c:\Users\HY\Desktop\8月业务员里程数_已填任务对应.xlsx")
BIZ_SRC = Path(r"c:\Users\HY\Desktop\cursor分析数据原始表.xlsx")
OUT_BACKUP = Path(r"d:\trae\crm-task\8月业务员里程数_已填任务对应.xlsx")

NAME_ALIASES = {"张桐?": "张桐菓", "王志刚1": "王志刚"}
BIZ_OVERRIDE = {"王金卫": "鸡窝"}

DB = dict(
    host="rr-uf614o6txbi981rfv0o.mysql.rds.aliyuncs.com",
    port=3306,
    user="pro_rd",
    password="O2KJIqmVQgbpZ1HO",
    charset="utf8mb4",
    cursorclass=DictCursor,
    connect_timeout=30,
    read_timeout=120,
)


def norm(n: str) -> str:
    n = (n or "").strip()
    return NAME_ALIASES.get(n, n)


def load_official_sales_biz():
    wb = openpyxl.load_workbook(BIZ_SRC, data_only=True)
    ws = wb["8月业务员创建任务和完成情况"]
    m = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        person = norm(str(r[3] or ""))
        biz = str(r[4] or "").strip()
        if person and biz in ("鸡窝", "大宗"):
            m[person] = biz
    wb.close()
    m.update(BIZ_OVERRIDE)
    return m


def classify_from_dept(dept: str | None) -> str | None:
    if not dept:
        return None
    d = dept
    # 优先级：采购 / 技术 > 业务线销售
    if "采购" in d:
        return "采购"
    if any(k in d for k in ("技术支持", "产品支持", "技术部", "技术老师", "饲料项目组")):
        return "技术老师"
    if "鸡窝" in d:
        return "鸡窝"
    if "大宗" in d:
        return "大宗"
    # 站台名常见于鸡窝销售编制，但无「鸡窝」字样时不武断
    return None


def load_salary_biz(names: list[str]):
    """取每人薪资部门（从近到远）推断业务类型。"""
    if not names:
        return {}, {}
    conn = pymysql.connect(**DB)
    cur = conn.cursor()
    rows_by = defaultdict(list)
    for i in range(0, len(names), 80):
        chunk = names[i : i + 80]
        cur.execute(
            """
            SELECT sys_name, department_name, salary_month
            FROM liangbaba_performance.staff_salary_detail
            WHERE sys_name IN (%s) AND (is_del=0 OR is_del IS NULL)
            ORDER BY salary_month DESC
            """
            % ",".join(["%s"] * len(chunk)),
            chunk,
        )
        for r in cur.fetchall():
            rows_by[norm(r["sys_name"])].append(r)
    conn.close()

    result = {}
    debug = {}
    for name, rows in rows_by.items():
        # 近→远：优先命中采购/技术，再业务线
        picked = None
        for r in rows:
            t = classify_from_dept(r["department_name"])
            if t in ("采购", "技术老师"):
                picked = (t, r["department_name"], r["salary_month"])
                break
        if not picked:
            for r in rows:
                t = classify_from_dept(r["department_name"])
                if t:
                    picked = (t, r["department_name"], r["salary_month"])
                    break
        if picked:
            result[name] = picked[0]
            debug[name] = f"{picked[1]}@{picked[2]}"
    return result, debug


def load_purchase_names() -> set[str]:
    conn = pymysql.connect(**DB)
    cur = conn.cursor()
    names = set()
    for t, col in (
        ("corn_commission_solidify_detail", "purchase_name"),
        ("soybean_commission_solidify_detail", "purchase_name"),
    ):
        # 近一年采购提成出现即视为采购岗
        cur.execute(
            f"""
            SELECT DISTINCT {col} AS n
            FROM liangbaba_performance.{t}
            WHERE commission_month >= %s AND {col} IS NOT NULL AND {col}<>''
            """,
            ("2025-01",),
        )
        names.update(norm(r["n"]) for r in cur.fetchall())
    cur.execute(
        """
        SELECT DISTINCT purchase_sys_name AS n
        FROM liangbaba_performance.delivery_solidify_detail
        WHERE delivery_month >= %s AND purchase_sys_name IS NOT NULL AND purchase_sys_name<>''
        """,
        ("2025-01",),
    )
    names.update(norm(r["n"]) for r in cur.fetchall() if r["n"])
    conn.close()
    return names


def main():
    official = load_official_sales_biz()
    src = MILE
    try:
        wb = openpyxl.load_workbook(src)
    except PermissionError:
        src = OUT_BACKUP
        wb = openpyxl.load_workbook(src)
    ws = wb[wb.sheetnames[0]]
    headers = [c.value for c in ws[1]]
    people = sorted(
        {
            norm(str(ws.cell(r, 8).value or ""))
            for r in range(2, ws.max_row + 1)
            if ws.cell(r, 8).value
        }
    )
    salary_biz, salary_debug = load_salary_biz(people)
    purchase_names = load_purchase_names()

    MANUAL = {"梁东君": "技术老师"}

    type_map = {}
    source = {}
    for p in people:
        if p in official:
            type_map[p] = official[p]
            source[p] = "任务业务字段"
        elif p in MANUAL:
            type_map[p] = MANUAL[p]
            source[p] = "手工"
        elif p in salary_biz:
            type_map[p] = salary_biz[p]
            source[p] = f"薪资部门:{salary_debug.get(p, '')}"
        elif p in purchase_names:
            type_map[p] = "采购"
            source[p] = "采购提成/提货"
        else:
            type_map[p] = None
            source[p] = "未识别"

    if "业务类型" in headers:
        col = headers.index("业务类型") + 1
    else:
        col = len(headers) + 1
        ws.cell(1, col, "业务类型")

    stats = Counter()
    unknown = []
    for r in range(2, ws.max_row + 1):
        name = norm(str(ws.cell(r, 8).value or ""))
        val = type_map.get(name)
        ws.cell(r, col, val)
        stats[val if val else "空"] += 1
        if not val and name and name not in unknown:
            unknown.append(name)

    outs = [MILE, OUT_BACKUP, Path(r"c:\Users\HY\Desktop\8月业务员里程数_已填任务对应_业务类型.xlsx")]
    for out in outs:
        try:
            wb.save(out)
            print("saved", out)
        except PermissionError:
            print("skip locked", out)

    print("stats", dict(stats))
    print("people", Counter(type_map.values()))
    print("unknown", unknown)
    for p in unknown:
        print(" ", p, source.get(p), salary_debug.get(p))
    print("\n非任务业务员:")
    for p in people:
        if p not in official:
            print(f"  {p}: {type_map.get(p)} | {source.get(p)}")


if __name__ == "__main__":
    main()
