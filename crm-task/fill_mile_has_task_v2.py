# -*- coding: utf-8 -*-
"""按正式「业务」字段重填「是否有对应的任务」。

- 鸡窝：仅 8.24–8.31；窗口外留空
- 大宗：8.1–8.31
- 非鸡窝/大宗（技术老师、采购等）：整列留空
- 王金卫：业务=鸡窝（8月调岗）
- 不区分审核状态；张桐?↔张桐菓
"""
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime, date
from pathlib import Path

import openpyxl

MILE_PATHS = [
    Path(r"c:\Users\HY\Desktop\8月业务员里程数.xlsx"),
    Path(r"c:\Users\HY\Desktop\8月业务员里程数_已填任务对应.xlsx"),
]
TASK_PATH = Path(r"c:\Users\HY\Desktop\8月任务列数据.xlsx")
BIZ_SRC = Path(r"c:\Users\HY\Desktop\cursor分析数据原始表.xlsx")
OUT_DESKTOP = Path(r"c:\Users\HY\Desktop\8月业务员里程数_已填任务对应.xlsx")
OUT_BACKUP = Path(r"d:\trae\crm-task\8月业务员里程数_已填任务对应.xlsx")

NAME_ALIASES = {"张桐?": "张桐菓", "王志刚1": "王志刚"}
# 强制覆盖（调岗等）
BIZ_OVERRIDE = {"王金卫": "鸡窝"}

JIWO_START, JIWO_END = date(2026, 8, 24), date(2026, 8, 31)
DZ_START, DZ_END = date(2026, 8, 1), date(2026, 8, 31)


def norm(n: str) -> str:
    n = (n or "").strip()
    return NAME_ALIASES.get(n, n)


def parse_day(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()[:10]
    for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except Exception:
            continue
    return None


def parse_range(s):
    if s is None:
        return None, None
    text = str(s).strip()
    if " ~ " in text:
        a, b = text.split(" ~ ", 1)
    elif "~" in text:
        a, b = text.split("~", 1)
    else:
        d = parse_day(text)
        return d, d
    return parse_day(a.strip()), parse_day(b.strip())


def load_biz_map():
    wb = openpyxl.load_workbook(BIZ_SRC, data_only=True)
    ws = wb["8月业务员创建任务和完成情况"]
    m = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        person = norm(str(r[3] or ""))
        biz = str(r[4] or "").strip()
        if person and biz in ("鸡窝", "大宗"):
            m[person] = biz
    wb.close()
    m.update(BIZ_OVERRIDE)
    return m


def load_task_ranges():
    wb = openpyxl.load_workbook(TASK_PATH, data_only=True)
    ws = wb["0"]
    by_name = defaultdict(list)
    for r in ws.iter_rows(min_row=2, values_only=True):
        person = norm(str(r[3] or ""))
        start, end = parse_range(r[6])
        if not person or not start or not end:
            continue
        if end < start:
            start, end = end, start
        by_name[person].append((start, end))
    wb.close()
    return by_name


def window_for(biz: str | None):
    if biz == "鸡窝":
        return JIWO_START, JIWO_END
    if biz == "大宗":
        return DZ_START, DZ_END
    return None  # 技术老师/采购等：不填


def covered(ranges, day: date) -> bool:
    return any(s <= day <= e for s, e in ranges)


def fill(path: Path, biz_map, by_name):
    wb = openpyxl.load_workbook(path)
    ws = wb[wb.sheetnames[0]]
    headers = [c.value for c in ws[1]]
    try:
        col = headers.index("是否有对应的任务") + 1
    except ValueError:
        col = len(headers) + 1
        ws.cell(1, col, "是否有对应的任务")

    stats = Counter()
    wjw = []
    for r in range(2, ws.max_row + 1):
        name = norm(str(ws.cell(r, 8).value or ""))
        day = parse_day(ws.cell(r, 13).value)
        if not name or not day:
            ws.cell(r, col, None)
            stats["空_缺字段"] += 1
            continue

        biz = biz_map.get(name)
        win = window_for(biz)
        if win is None:
            ws.cell(r, col, None)
            stats["空_非业务员"] += 1
            continue

        w0, w1 = win
        if not (w0 <= day <= w1):
            ws.cell(r, col, None)
            stats["空_窗口外"] += 1
            continue

        val = "是" if covered(by_name.get(name, []), day) else "否"
        ws.cell(r, col, val)
        stats[val] += 1
        if name == "王金卫":
            wjw.append((day.isoformat(), val))

    return wb, stats, wjw


def main():
    biz_map = load_biz_map()
    by_name = load_task_ranges()
    print("biz", Counter(biz_map.values()), "王金卫", biz_map.get("王金卫"))
    print("task people", len(by_name))

    src = next((p for p in MILE_PATHS if p.exists()), None)
    if src is None:
        raise SystemExit("no mileage file")

    wb, stats, wjw = fill(src, biz_map, by_name)
    print("stats", dict(stats))
    print("王金卫 before 8.24", [x for x in wjw if x[0] < "2026-08-24"][:3], "count", sum(1 for x in wjw if x[0] < "2026-08-24"))
    print("王金卫 from 8.24", [x for x in wjw if x[0] >= "2026-08-24"])

    for out in (OUT_DESKTOP, OUT_BACKUP):
        try:
            wb.save(out)
            print("saved", out)
        except PermissionError:
            alt = out.with_name(out.stem + "_v2.xlsx")
            wb.save(alt)
            print("locked, saved", alt)


if __name__ == "__main__":
    main()
