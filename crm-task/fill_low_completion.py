# -*- coding: utf-8 -*-
"""填充「有任务但是完成度不达标」：

- 仅当「是否有对应的任务」=是（当天有任务覆盖）
- 大宗阈值 100% → 不达标写「未达标100%」
- 鸡窝阈值 75% → 不达标写「未达标75%」
- 覆盖任务取最高完成度；低于阈值且覆盖任务均无申诉 → 不达标
- 是否申诉：有值=有申诉（不算不达标），空=没申诉
"""
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime, date
from pathlib import Path

import openpyxl

MILE = Path(r"c:\Users\HY\Desktop\8月业务员里程数_已填任务对应.xlsx")
TASK = Path(r"c:\Users\HY\Desktop\8月任务列数据.xlsx")
OUT_BACKUP = Path(r"d:\trae\crm-task\8月业务员里程数_已填任务对应.xlsx")

NAME_ALIASES = {"张桐?": "张桐菓", "王志刚1": "王志刚"}
THRESH = {"大宗": 1.0, "鸡窝": 0.75}
LABEL = {"大宗": "未达标100%", "鸡窝": "未达标75%"}


def norm(n: str) -> str:
    n = (n or "").strip()
    return NAME_ALIASES.get(n, n)


def parse_day(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()[:10]
    for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except Exception:
            continue
    return None


def parse_range(s):
    if s is None:
        return None, None
    text = str(s).strip()
    if " ~ " in text:
        a, b = text.split(" ~ ", 1)
    elif "~" in text:
        a, b = text.split("~", 1)
    else:
        d = parse_day(text)
        return d, d
    return parse_day(a.strip()), parse_day(b.strip())


def parse_pct(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        x = float(v)
        return x / 100.0 if x > 1 else x
    try:
        return float(str(v).strip().replace("%", "")) / 100.0
    except Exception:
        return None


def has_appeal(v) -> bool:
    """是否申诉有值=有申诉，空=没申诉。"""
    if v is None:
        return False
    if isinstance(v, str) and v.strip() == "":
        return False
    return True


def load_tasks():
    wb = openpyxl.load_workbook(TASK, data_only=True)
    ws = wb["0"]
    by_name = defaultdict(list)
    for r in ws.iter_rows(min_row=2, values_only=True):
        person = norm(str(r[3] or ""))
        start, end = parse_range(r[6])
        if not person or not start or not end:
            continue
        if end < start:
            start, end = end, start
        pct = parse_pct(r[5])
        by_name[person].append(
            {
                "start": start,
                "end": end,
                "pct": pct if pct is not None else 0.0,
                "pct_raw": r[5],
                "appeal": has_appeal(r[8]),
                "status": str(r[2] or ""),
            }
        )
    wb.close()
    return by_name


def main():
    by_name = load_tasks()
    wb = openpyxl.load_workbook(MILE)
    ws = wb[wb.sheetnames[0]]
    headers = [c.value for c in ws[1]]

    col_name = headers.index("biz_user_name") + 1
    col_day = headers.index("expense_time") + 1
    col_biz = headers.index("业务类型") + 1
    col_has = headers.index("是否有对应的任务") + 1
    col_low = headers.index("有任务但是完成度不达标") + 1

    stats = Counter()
    samples = []

    for r in range(2, ws.max_row + 1):
        name = norm(str(ws.cell(r, col_name).value or ""))
        day = parse_day(ws.cell(r, col_day).value)
        biz = str(ws.cell(r, col_biz).value or "").strip()
        has_task = str(ws.cell(r, col_has).value or "").strip()

        # 仅「有对应任务」的鸡窝/大宗行才判断不达标
        if has_task != "是" or biz not in THRESH or not name or not day:
            ws.cell(r, col_low, None)
            stats["空"] += 1
            continue

        covering = [
            t
            for t in by_name.get(name, [])
            if t["start"] <= day <= t["end"]
        ]
        if not covering:
            # 与「是」不一致时保守留空
            ws.cell(r, col_low, None)
            stats["空_无覆盖异常"] += 1
            continue

        if any(t["appeal"] for t in covering):
            ws.cell(r, col_low, None)
            stats["空_有申诉"] += 1
            continue

        best_pct = max(t["pct"] for t in covering)
        thr = THRESH[biz]
        if best_pct + 1e-9 < thr:
            val = LABEL[biz]
            ws.cell(r, col_low, val)
            stats[val] += 1
            if len(samples) < 8:
                samples.append((name, biz, day.isoformat(), best_pct, val))
        else:
            ws.cell(r, col_low, None)
            stats["空_已达标"] += 1

    saved = []
    for out in (MILE, OUT_BACKUP):
        try:
            wb.save(out)
            saved.append(str(out))
        except PermissionError:
            alt = out.with_name(out.stem + "_完成度不达标.xlsx")
            wb.save(alt)
            saved.append(str(alt) + " (原文件占用)")

    print("stats", dict(stats))
    print("samples", samples)
    print("saved", saved)


if __name__ == "__main__":
    main()
